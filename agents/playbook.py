import os
import sys
from datetime import date
from docx import Document as DocxDocument
from docx.shared import Pt
from services.ai_core import chat_completion
from services.hana import get_connection

_SYSTEM = (
    "Sei un knowledge manager di una practice di consulenza. "
    "Genera un playbook anonimizzato in italiano. "
    "Rimuovi TUTTI i nomi di persone, nomi del cliente, email e riferimenti identificativi. "
    "Mantieni la conoscenza metodologica e tecnica."
)

_PROMPT = (
    "Analizza la memoria del progetto e genera un playbook strutturato con queste 7 sezioni:\n\n"
    "1. CONTESTO E OBIETTIVI — settore, dimensione, durata (senza nominare il cliente)\n"
    "2. APPROCCIO METODOLOGICO — fasi, framework, governance adottati\n"
    "3. DECISIONI CHIAVE E MOTIVAZIONI — le 5-8 più importanti con il perché\n"
    "4. PROBLEMI INCONTRATI E SOLUZIONI — formato: Problema → Soluzione → Risultato\n"
    "5. PATTERN E LESSONS LEARNED — 3-5 pattern riutilizzabili in altri progetti\n"
    "6. RACCOMANDAZIONI — 4-6 raccomandazioni pratiche per progetti simili\n"
    "7. STACK TECNOLOGICO — architettura e configurazioni chiave (senza dati sensibili)\n\n"
    "MEMORIA DEL PROGETTO:\n{memoria}"
)

_MAX_MEMORY = 20000


class PlaybookAgent:
    def genera(self, project_id: str = None, output_path: str = "playbook.docx") -> str:
        if project_id is None:
            project_id = os.getenv("PROJECT_ID", "demo")

        conn = get_connection()
        cursor = conn.cursor()
        memoria_parts = []

        try:
            cursor.execute(
                "SELECT SOURCE_NAME, CONTENT FROM IPMP_DOCS "
                "WHERE PROJECT_ID=? AND SOURCE_NAME NOT LIKE 'CONTRATTO::%' "
                "ORDER BY CREATED_AT",
                [project_id],
            )
            for row in cursor.fetchall():
                memoria_parts.append(f"[{row[0]}] {row[1]}")

            cursor.execute(
                "SELECT DESCRIZIONE, AUTORE, DATA_DECISIONE, MOTIVAZIONE FROM IPMP_DECISIONS "
                "WHERE PROJECT_ID=? ORDER BY CREATED_AT",
                [project_id],
            )
            for row in cursor.fetchall():
                memoria_parts.append(
                    f"[DECISIONE] {row[0]} (Autore: {row[1]}, Data: {row[2]}) — {row[3]}"
                )
        finally:
            cursor.close()
            conn.close()

        memoria = "\n\n".join(memoria_parts)
        if len(memoria) > _MAX_MEMORY:
            memoria = memoria[:_MAX_MEMORY]

        try:
            playbook_text = chat_completion(
                messages=[{"role": "user", "content": _PROMPT.format(memoria=memoria)}],
                system=_SYSTEM,
                temperature=0.3,
            )
        except Exception as e:
            print(f"[PlaybookAgent] Errore LLM: {e}", file=sys.stderr)
            raise

        # Costruisce il .docx
        doc = DocxDocument()
        doc.add_heading(f"Project Playbook — {date.today().isoformat()}", 0)
        doc.add_paragraph(f"Progetto: {project_id} | Generato da i-PMP")
        doc.add_paragraph("")

        current_section = None
        for line in playbook_text.split("\n"):
            stripped = line.strip()
            if not stripped:
                continue
            # Rileva intestazioni di sezione (numerate o in maiuscolo)
            if stripped and (
                stripped[0].isdigit() and "." in stripped[:3]
                or stripped.isupper() and len(stripped) > 3
            ):
                doc.add_heading(stripped, level=1)
                current_section = stripped
            elif stripped.startswith("•") or stripped.startswith("-"):
                doc.add_paragraph(stripped, style="List Bullet")
            else:
                doc.add_paragraph(stripped)

        doc.save(output_path)
        return output_path
